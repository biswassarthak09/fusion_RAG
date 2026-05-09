# template_filler.py
#
# Fills Template_01_Excel_Massnahmenliste by:
#   1. Discovering all measures mentioned across the 5 PDF knowledge-base docs
#   2. For each measure, running targeted RAG queries to fill each column
#   3. Writing a new row per measure and saving a timestamped output file
#
# Column mapping (Template_01):
#   A: ID              — auto-incremented (M-001, M-002, ...)
#   B: Maßnahme Titel  — [FILL: Wie lautet der Titel / Name dieser Maßnahme?]
#   C: Kategorie       — [FILL: Welcher Kategorie gehört diese Maßnahme an?]
#   D: Kurzbeschreibung— [FILL: Kurzbeschreibung dieser Maßnahme?]
#   E: Priorität       — [FILL: Welche Priorität hat diese Maßnahme (hoch/mittel/niedrig)?]
#   F: Startjahr       — [FILL: In welchem Jahr soll diese Maßnahme beginnen?]
#   G: Endjahr         — [FILL: In welchem Jahr soll diese Maßnahme enden?]
#   H: Verantwortlich  — [FILL: Wer ist verantwortlich für diese Maßnahme?]
#   I: Abhängigkeiten  — [FILL: Welche Abhängigkeiten oder Risiken hat diese Maßnahme?]
#   J: Quelle          — filled automatically from chunk metadata (filename)
#   K: Fundstelle      — filled automatically from chunk metadata (page)
#   L: Unsicherheit    — [FILL: Gibt es Unsicherheiten bei dieser Maßnahme (Ja/Nein)?]

import os
import re
import shutil
from datetime import datetime

import openpyxl

from langchain_ollama import ChatOllama
from langchain_core.prompts import PromptTemplate
from langchain_core.output_parsers import StrOutputParser
from vector_db import get_specific_doc_retriever, get_keyword_chunks

# ── Config ─────────────────────────────────────────────────────────────────────
TEMPLATE_PATH = "./templates/Template_01_Excel_Massnahmenliste_Waermeplanung_SHARED_FINAL.xlsx"
OUTPUT_DIR    = "./output"
HEADER_ROW    = 1   # row index (1-based) of the column headers
DATA_START_ROW = 2  # first data row — we overwrite from here
SESSION_ID    = "global"

# Column indices (1-based, matching the template)
COL_ID          = 1
COL_TITEL       = 2
COL_KATEGORIE   = 3
COL_BESCHREIBUNG= 4
COL_PRIORITAET  = 5
COL_STARTJAHR   = 6
COL_ENDJAHR     = 7
COL_VERANTWORTL = 8
COL_ABHAENGIG   = 9
COL_QUELLE      = 10
COL_FUNDSTELLE  = 11
COL_UNSICHERHEIT= 12

# LLM (same model as the main pipeline)
llm = ChatOllama(model="llama3.2", temperature=0)

# ── Helpers ────────────────────────────────────────────────────────────────────

def _rag_query(question: str, vectorstore, source_file: str) -> tuple[str, str, str]:
    """
    Runs a single RAG query against one source file.
    Returns (answer, source_basename, page_str).
    Uses the keyword top-up retriever from the main pipeline.
    """
    retriever = get_specific_doc_retriever(vectorstore, source_file, SESSION_ID)
    docs = retriever.invoke(question)

    # Keyword top-up — reuse same logic as nodes.py
    keyword_terms = [
        "maßnahme", "maßnahmen", "ziel", "aktion", "priorität",
        "verantwortlich", "startjahr", "endjahr", "umsetzung",
        "verpflichtet", "pflicht", "konzessionsabgabe"
    ]
    kw_docs = get_keyword_chunks(vectorstore, source_file, keyword_terms)
    seen = {d.page_content[:80] for d in docs}
    for d in kw_docs:
        if d.page_content[:80] not in seen:
            docs.append(d)
            seen.add(d.page_content[:80])

    if not docs:
        return "Unklar", os.path.basename(source_file), "?"

    # Sort: keyword-matched first, cap at 6
    def _has_kw(d):
        return any(k in d.page_content.lower() for k in keyword_terms)
    priority = [d for d in docs if _has_kw(d)]
    rest     = [d for d in docs if not _has_kw(d)]
    top_docs = (priority + rest)[:6]

    # Build context
    context = ""
    for doc in top_docs:
        src  = os.path.basename(doc.metadata.get("source", "?"))
        page = doc.metadata.get("page", "?")
        context += f"--- QUELLE: {src} | SEITE: {page} ---\n{doc.page_content}\n\n"

    # Best source / page from top chunk
    best_source = os.path.basename(top_docs[0].metadata.get("source", source_file))
    best_page   = str(top_docs[0].metadata.get("page", "?"))

    template_str = """Du bist ein präziser Assistent. Beantworte die Frage kurz und direkt, nur auf Basis des Kontexts.
Wenn die Information nicht im Kontext steht, antworte mit: Unklar

Kontext:
{context}

Frage: {question}
Antwort (max. 2 Sätze):"""

    prompt  = PromptTemplate.from_template(template_str)
    chain   = prompt | llm | StrOutputParser()
    answer  = chain.invoke({"context": context, "question": question}).strip()

    return answer, best_source, best_page


def _discover_measures(vectorstore, source_file: str) -> list[str]:
    """
    Asks the LLM to list all concrete measures mentioned in the document.
    Returns a list of measure title strings.
    """
    print(f"  🔎 Discovering measures in: {os.path.basename(source_file)}")

    # Fetch all chunks from this file and take the first 12 (broad overview)
    all_chunks = vectorstore.get(where={"source": source_file})
    if not all_chunks["documents"]:
        return []

    # Prioritise chunks that mention "Maßnahme" / "Maßnahmen"
    measure_kw = ["maßnahme", "maßnahmen", "aktion", "umsetzung", "ziel"]
    priority, rest = [], []
    for text in all_chunks["documents"]:
        if any(k in text.lower() for k in measure_kw):
            priority.append(text)
        else:
            rest.append(text)

    selected = (priority + rest)[:12]
    context  = "\n\n---\n\n".join(selected)

    template_str = """Du bist ein Analyst. Lies den folgenden Text und liste alle konkreten Maßnahmen, Aktionen oder Projekte auf, die im Text erwähnt werden.

Regeln:
- Gib nur Maßnahmen aus, die explizit im Text stehen.
- Jede Maßnahme auf einer eigenen Zeile, beginnend mit "- "
- Maximal 15 Maßnahmen.
- Wenn keine Maßnahmen erkennbar sind, schreibe: KEINE

Text:
{context}

Maßnahmen:"""

    prompt  = PromptTemplate.from_template(template_str)
    chain   = prompt | llm | StrOutputParser()
    result  = chain.invoke({"context": context}).strip()

    if "keine" in result.lower() or not result:
        return []

    measures = []
    for line in result.splitlines():
        line = line.strip().lstrip("-•*").strip()
        if len(line) > 5:
            measures.append(line)

    return measures


def _fill_row_for_measure(
    measure_title: str,
    source_file: str,
    vectorstore,
    row_id: int
) -> dict:
    """
    Runs one RAG query per column for a given measure and returns a dict
    keyed by column index.
    """
    base = os.path.basename(source_file)
    print(f"    📝 Filling columns for: {measure_title[:60]}...")

    # Column queries — each question references the measure by name
    queries = {
        COL_KATEGORIE:   f"Welcher Kategorie gehört die Maßnahme '{measure_title}' an? (z.B. Wärmenetz, Gebäude, Quartier, Governance, Netz)",
        COL_BESCHREIBUNG:f"Gib eine kurze Beschreibung der Maßnahme '{measure_title}'.",
        COL_PRIORITAET:  f"Welche Priorität hat die Maßnahme '{measure_title}'? (hoch, mittel, niedrig oder Unklar)",
        COL_STARTJAHR:   f"In welchem Jahr soll die Maßnahme '{measure_title}' beginnen? Nur die Jahreszahl oder 'Unklar'.",
        COL_ENDJAHR:     f"In welchem Jahr soll die Maßnahme '{measure_title}' abgeschlossen sein? Nur die Jahreszahl oder 'Unklar'.",
        COL_VERANTWORTL: f"Wer ist verantwortlich für die Umsetzung der Maßnahme '{measure_title}'?",
        COL_ABHAENGIG:   f"Welche Abhängigkeiten, Voraussetzungen oder Risiken hat die Maßnahme '{measure_title}'?",
        COL_UNSICHERHEIT:f"Gibt es Unsicherheiten bei der Maßnahme '{measure_title}'? Antworte mit Ja oder Nein.",
    }

    row = {
        COL_ID:    f"M-{row_id:03d}",
        COL_TITEL: measure_title,
    }

    source_for_row = base
    page_for_row   = "?"

    for col, question in queries.items():
        answer, src, page = _rag_query(question, vectorstore, source_file)
        row[col] = answer
        if col == COL_KATEGORIE:      # capture metadata from first substantive query
            source_for_row = src
            page_for_row   = page

    row[COL_QUELLE]     = source_for_row
    row[COL_FUNDSTELLE] = page_for_row

    return row


# ── Main entry point ──────────────────────────────────────────────────────────

def fill_template(vectorstore, pdf_sources: list[str]) -> str:
    """
    Main function called from main_bfh.py (option 6).
    Discovers measures across all PDF sources, fills a row per measure,
    and saves a timestamped copy of the template.

    Returns the path of the saved output file.
    """
    os.makedirs(OUTPUT_DIR, exist_ok=True)

    # Only keep PDF sources (filter out URLs)
    pdf_only = [s for s in pdf_sources if s.startswith("./knowledge_base/")]
    print(f"\n📂 Processing {len(pdf_only)} PDF document(s)...\n")

    # ── Step 1: Discover all measures across all PDFs ──────────────────────────
    all_measures: list[tuple[str, str]] = []  # (measure_title, source_file)
    for source_file in pdf_only:
        measures = _discover_measures(vectorstore, source_file)
        print(f"     → Found {len(measures)} measure(s)")
        for m in measures:
            all_measures.append((m, source_file))

    if not all_measures:
        print("⚠️  No measures found across the documents.")
        return ""

    print(f"\n✅ Total measures discovered: {len(all_measures)}")

    # ── Step 2: Load template and prepare output workbook ─────────────────────
    out_name = f"Template_01_FILLED_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    out_path = os.path.join(OUTPUT_DIR, out_name)
    shutil.copy2(TEMPLATE_PATH, out_path)

    wb = openpyxl.load_workbook(out_path)
    ws = wb["Maßnahmenliste"]

    # Clear the placeholder example row (row 2) — we'll write real data from row 2
    for col in range(1, 13):
        ws.cell(row=DATA_START_ROW, column=col).value = None

    # ── Step 3: Fill one row per measure ──────────────────────────────────────
    print(f"\n🔧 Filling {len(all_measures)} row(s)...\n")
    for row_idx, (measure_title, source_file) in enumerate(all_measures, start=1):
        print(f"\n  [{row_idx}/{len(all_measures)}] {os.path.basename(source_file)}")
        row_data = _fill_row_for_measure(measure_title, source_file, vectorstore, row_idx)

        excel_row = DATA_START_ROW + row_idx - 1
        for col, value in row_data.items():
            ws.cell(row=excel_row, column=col).value = value

    wb.save(out_path)
    print(f"\n✅ Template filled and saved to: {out_path}")
    return out_path
